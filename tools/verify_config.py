"""Compare the minimal Broad tables with the owner's pin workbook."""
from __future__ import annotations

import argparse
import re
from pathlib import Path

import openpyxl


ROOT = Path(__file__).resolve().parents[1]


def verify(sheet: Path) -> None:
    ws = openpyxl.load_workbook(sheet, read_only=True, data_only=True).active
    source = (ROOT / "Broad/src/DrvCfg_EP15K.c").read_text(encoding="utf-8")
    gpio_part = source.split("ST_GPIO_CFG_ITEM stGpioCfgItem[] =", 1)[1].split("/* stGpioCfg */", 1)[0]
    pins: dict[int, tuple[str, str, str, str, str]] = {}
    for match in re.finditer(
        r"\{\s*(\d+),\s*(HW_GPIO_\w+),\s*(HW_DIR_\w+),\s*"
        r"HW_GPIO_\w+,\s*(HW_GPIO_PUD_\w+),\s*(HW_GPIO_OD_\w+),\s*"
        r"(HW_GPIO_QUAL_\w+)\s*\}", gpio_part
    ):
        pin = int(match[1])
        assert pin not in pins, f"duplicate GPIO{pin}"
        pins[pin] = match.groups()[1:]

    checks = 0
    for row in range(33, 57):
        name, pin_text = ws.cell(row, 1).value, ws.cell(row, 4).value
        if not re.fullmatch(r"EPWM\d+[AB]", str(name)):
            continue
        pin = int(str(pin_text).removeprefix("GPIO"))
        mux, direction, *_ = pins[pin]
        assert direction == "HW_DIR_DO", f"{name} GPIO{pin} direction"
        expected_mux = "HW_GPIO_MUX_PWM" if pin < 16 else "HW_GPIO_FUNC_PERIPH5"
        assert mux == expected_mux, f"{name} GPIO{pin} mux {mux}"
        checks += 1

    for row, direction in [(r, "HW_DIR_DI") for r in range(58, 75)] + [
        (r, "HW_DIR_DO") for r in range(76, 100)
    ]:
        name, pin_text = ws.cell(row, 2).value, ws.cell(row, 4).value
        if not isinstance(pin_text, str) or not pin_text.startswith("GPIO") or name == "/":
            continue
        pin = int(pin_text[4:])
        assert pin in pins, f"missing GPIO{pin} for {name}"
        assert pins[pin][0] == "HW_GPIO_MUX_GPIO" and pins[pin][1] == direction, f"GPIO{pin} mode"
        checks += 1

    for row, mux, direction in [
        (101, "CAN", "DI"), (102, "CAN", "DO"),
        (103, "CAN", "DO"), (104, "CAN", "DI"),
        (106, "SCI", "DO"), (107, "SCI", "DI"),
        (108, "SCI", "DO"), (109, "SCI", "DI"),
        (110, "SCI", "DO"), (111, "SCI", "DI"),
    ]:
        pin = int(ws.cell(row, 4).value[4:])
        assert pins[pin][:2] == (f"HW_GPIO_MUX_{mux}", f"HW_DIR_{direction}"), f"GPIO{pin} {mux}"
        checks += 1

    adc_part = source.split("const ST_ADC_CH_CFG_ITEM stAdcChCfgItem[] =", 1)[1].split("/* stAdcCfgItem */", 1)[0]
    entries = [tuple(int(x) for x in m) for m in re.findall(r"\{\s*(\d+),\s*(\d+),\s*(\d+)\s*\}", adc_part)]
    assert len(entries) == 64, f"ADC SOC table: {len(entries)} entries"
    configured = {f"{'ABCD'[i % 4]}{ch}" for i, (trig, acq, ch) in enumerate(entries) if trig != 0}
    expected = {str(ws.cell(r, 1).value) for r in range(2, 20)} | {"D14", "D15"}
    assert configured == expected, f"ADC missing={expected-configured} extra={configured-expected}"
    assert entries[15 * 4] == (5, 120, 0), "ADCA SOC15 anchor"
    assert all(entries[soc * 4 + adc] == (0, 0, 0) for soc in (8, 9) for adc in range(4))
    checks += len(expected) + 2

    for pin in (13, 17, 40, 41, 77, 80, 81, 82):
        assert pins[pin][2:] == ("HW_GPIO_PUD_ENABLE", "HW_GPIO_OD_DISABLE", "HW_GPIO_QUAL_ASYNC")
        checks += 1
    for pin in (36, 37, 38, 39, 48, 49, 89, 90, 93, 94):
        assert pins[pin][3] == "HW_GPIO_OD_DISABLE", f"GPIO{pin} open drain"
        checks += 1

    pwm_part = source.split("ST_PWM_CFG_ITEM stPwmCfgItem[] =", 1)[1].split("/* stPwmCfg */", 1)[0]
    modules = re.findall(r"HW_PWM_ID(\d+),\s*HW_PWM_TYPE_(INV|PV|BUS|FAN),\s*(\d+)U.*?TB_COUNT_(UPDOWN|UP),", pwm_part)
    assert len(modules) == 12 and {int(m[0]) for m in modules} == set(range(12))
    for pwm_id, kind, period, mode in modules:
        hz = 100_000_000 // (int(period) * 2 if mode == "UPDOWN" else int(period) + 1)
        assert hz == (10_000 if kind == "FAN" else 20_000), f"ePWM{int(pwm_id)+1}={hz}Hz"
        checks += 1
    syncs = re.findall(
        r"HW_PWM_ID(\d+),\s*HW_PWM_TYPE_\w+,.*?TB_COUNT_\w+,\s*"
        r"HW_PWM_\w+,\s*\d+U,\s*HW_EPWM_SYNC_IN_FROM_\w+,\s*"
        r"HW_EPWM_SYNC_OUT_FROM_(\w+)", pwm_part
    )
    assert len(syncs) == 12
    sync_by_id = {int(pwm_id): mode for pwm_id, mode in syncs}
    for pwm_id, expected_mode in {
        0: "CNTR_ZERO", 1: "SYNCIN", 3: "SYNCIN", 4: "SYNCIN",
        6: "CNTR_ZERO", 7: "SYNCIN", 9: "SYNCIN",
    }.items():
        assert sync_by_id[pwm_id] == expected_mode, f"ePWM{pwm_id+1} sync chain"
        checks += 1
    fan = pwm_part.split("HW_PWM_ID11, HW_PWM_TYPE_FAN", 1)[1]
    assert re.search(r"\{AQ_NO_ACTION,\s*AQ_NO_ACTION,\s*AQ_SET,\s*AQ_CLEAR", fan), "EPWM12B action"
    checks += 1
    assert "{HW_SCI_ID_A,  19200UL" in source
    assert "{HW_SCI_ID_C, 115200UL" in source
    assert "{HW_SCI_ID_D,  38400UL" in source
    assert "{HW_CAN_ID_A, 500000UL" in source
    assert "{HW_CAN_ID_B, 1000000UL" in source
    print(f"PASS: {checks + 5} pin, ADC, frequency and communication checks")


if __name__ == "__main__":
    parser = argparse.ArgumentParser()
    parser.add_argument("--pin-sheet", type=Path, required=True)
    verify(parser.parse_args().pin_sheet)
